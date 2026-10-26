"""Source-preserving 2024 Panini Three Count, Select and PhotoGenic WWE import.

Mirrored Panini-format workbooks are retained, with independent workbook comparison.
No absent base versions or parallel applicability are synthesized.
"""
import argparse, hashlib, json, re
from collections import OrderedDict, Counter
from pathlib import Path
from catalogue_io import load_catalogue, save_catalogue
ROOT=Path(__file__).resolve().parents[1]
SOURCE_DIR=ROOT/'sources/wwe-2024-batch'
BASELINE='181191c8139e31b7c5913cf7ec1dafae8aebd4e5'
CONFIGS=OrderedDict([
 ('three-count',dict(id='2024-panini-three-count-wwe',prefix='2024-P3CWWE',name='2024 Panini Three Count WWE',short='Three Count WWE 2024',date='2024-12-06',month='11',accent='#bf9a64',rows=3333,
  families=[('Attitude Era Icons Signatures','AEIS','Autograph',17),('Base','BASE','Base',100),('Base Legends','BL','Base',50),('Base Timeless Moments','BTM','Base',50),('Downtown','DT','Insert',20),('Jumbo Memorabilia Autographs','JMA','Autograph Relic',10),('Legendary Autographs','LA','Autograph',30),('Massive Mat Autographs','MMA','Autograph Relic',19),('Memorabilia Autographs','MA','Autograph Relic',40),('Prime Memorabilia Autographs','PMA','Autograph Relic',30),('Superstar Swatches Autographs','SSA','Autograph Relic',20),('Timeless Moments Autographs Horizontal','TMAH','Autograph',15),('Timeless Moments Autographs Vertical','TMAV','Autograph',14)])),
 ('select',dict(id='2024-panini-select-wwe',prefix='2024-PSWWE',name='2024 Panini Select WWE',short='Select WWE 2024',date='2024-11-08',month='10',accent='#bd995e',rows=11234,
  families=[('And NEW!','AN','Insert',20),('Artistic Selections','AS','Insert',5),('Autographed Memorabilia','AM','Autograph Relic',26),('Base Set - Concourse','CON','Base',100),('Base Set - Premier Level','PRE','Base',100),('Base Set - Ringside','RING','Base',100),('Color Wheel','CW','Insert',10),('Firestorm','FS','Insert',10),('Icons','IC','Insert',25),('Legendary Signatures','LS','Autograph',20),('NXT Alumni','NXT','Insert',25),('Phenomenon','PH','Insert',30),('Ringside Action Signatures','RAS','Autograph',30),('Select Future','SF','Insert',25),('Select Throwback','ST','Insert',25),('Selective Swatches','SSW','Relic',34),('Sensations','SEN','Insert',20),('Signature Selections','SLS','Autograph',94),('Signatures','SIG','Autograph',25),('Sparks','SP','Relic',35),('Turbo Charged','TC','Insert',30)])),
 ('photogenic',dict(id='2024-panini-photogenic-wwe',prefix='2024-PPGWWE',name='2024 Panini PhotoGenic WWE',short='PhotoGenic WWE 2024',date='2024-09-25',month='09',accent='#9c9fd2',rows=2924,
  families=[('Base Horizontal','BH','Base',100),('Base Vertical','BV','Base',101),('Championship Feels','CF','Insert',10),('Down The Ramp','DTR','Insert',15),('Dual Autographs','DA','Autograph',10),('Immortalized','IM','Insert',15),('Photo Finish','PF','Insert',15),('Picture Perfect Signatures','PPS','Autograph',37),('Progressions','PRO','Insert',15),('Step By Step','SBS','Insert',15),('Superstar Signatures','SS','Autograph',47),('Zero Gravity','ZG','Insert',15)])),
])
for key,cfg in CONFIGS.items():
 title={'three-count':'Three-Count','select':'Select','photogenic':'PhotoGenic'}[key]
 cfg['source']=f"https://xcdn.checklistinsider.com/public/2024/{cfg['month']}/2024-Panini-{title}-WWE-Wrestling-Checklist-Downloads-Excel-spreadsheet.xlsx"
 cfg['guide']='https://www.checklistinsider.com/2024-panini-'+key+'-wwe'+('-wrestling' if key=='select' else '')
 cfg['center']='https://www.checklistcenter.com/2024-panini-'+key+'-wwe-wrestling-card-checklist/'

def clean(x):
 if isinstance(x,(int,float)) and not isinstance(x,bool) and int(x)==x:return str(int(x))
 return re.sub(r'\s+',' ',str(x or '')).strip()
def slug(x):return re.sub('[^A-Z0-9]+','-',x.upper()).strip('-')
def subjects(x):return sorted(clean(s) for s in x.split('/'))
def rowkey(r):
 people=subjects(r[2])
 if r[0].startswith('Progressions') and r[1]=='6':people=[p.replace('"The Fiend" Bray Wyatt','Bray Wyatt') for p in people]
 return (r[0],r[1],tuple(sorted(set(people))),tuple(sorted(set(subjects(r[3])))),r[4])
def read_source(key):return json.loads((SOURCE_DIR/(key+'.json')).read_text())

def extract(key):
 import openpyxl
 cfg=CONFIGS[key];p=SOURCE_DIR/(key+'-insider.xlsx');w=openpyxl.load_workbook(p,data_only=True)
 if w.active.max_column!=5 or list(next(w.active.values))!=['CARD SET','CARD #','ATHLETE','TEAM','SEQ']:raise ValueError('Unexpected source columns')
 rows=[[clean(x) for x in r] for r in list(w.active.values)[1:]]
 if len(rows)!=cfg['rows'] or any(not all(r[:3]) for r in rows):raise ValueError('Source count/required fields')
 for r in rows:
  if r[4] and (not r[4].isdigit() or int(r[4])<1):raise ValueError('Bad printed sequence')
 alt=SOURCE_DIR/('photogenic.xlsx' if key=='photogenic' else key+'-center.xlsx')
 ar=[[clean(x) for x in r[:5]] for r in list(openpyxl.load_workbook(alt,data_only=True).active.values)[1:] if r[0] and r[1]]
 # The second PhotoGenic workbook retains separate signer rows for duals.
 combined=OrderedDict()
 for r in ar:
  if key=='photogenic' and r[0].startswith('Dual Autographs'):
   g=combined.setdefault((r[0],r[1],r[4]),[r[0],r[1],[],[],r[4]]);g[2].append(r[2]);g[3].append(r[3])
  else:combined[('single',len(combined))]=r
 normalized=[]
 for r in combined.values():
  if isinstance(r[2],list):r=[r[0],r[1],'/'.join(r[2]),'/'.join(r[3]),r[4]]
  normalized.append(r)
 a,b=Counter(map(rowkey,rows)),Counter(map(rowkey,normalized))
 if a!=b:raise ValueError(f'Independent workbook mismatch {key}: primary-only {list((a-b).items())[:5]}, secondary-only {list((b-a).items())[:5]}')
 origin={}
 if key in ('three-count','select'):
  import xlrd
  op=SOURCE_DIR/(key+'.xls');sheet=xlrd.open_workbook(op).sheet_by_index(0)
  if sheet.row_values(2)!=['CARD #','CARD SET','ATHLETE','TEAM','SEQUENCE']:raise ValueError('Distributor columns changed')
  original=[[clean(sheet.cell_value(i,j)) for j in (1,0,2,3,4)] for i in range(3,sheet.nrows)]
  # Both original exports have the same single mojibake string for Seth's nickname.
  bad='Seth â€œFreakinâ€\x9d Rollins';fixed='Seth “Freakin” Rollins'
  repaired=[[cell.replace(bad,fixed) for cell in row] for row in original]
  if Counter(map(rowkey,repaired))!=a:raise ValueError('Distributor workbook does not reconcile')
  origin=dict(distributorUrl=f"https://gogts.net/wp-content/uploads/2024/{cfg['month']}/2024-Panini-{'Three-Count' if key=='three-count' else 'Select'}-WWE-Trading-Cards-Checklist.xls",distributorSha256=hashlib.sha256(op.read_bytes()).hexdigest(),distributorRows=len(original),encodingRepairs=sum(x!=y for x,y in zip(original,repaired)),distributorRowsOriginal=original)
 return dict(**origin,sourceUrl=cfg['source'],sourceSha256=hashlib.sha256(p.read_bytes()).hexdigest(),comparisonUrl=cfg['center'],comparisonSha256=hashlib.sha256(alt.read_bytes()).hexdigest(),comparisonRawRows=len(ar),rows=rows)

def classify(key,label):
 families=CONFIGS[key]['families']
 choices=[f for f in families if label==f[0] or label.startswith(f[0]+' ')]
 if not choices:raise ValueError('Unknown family '+label)
 family=max(choices,key=lambda f:len(f[0]));parallel=label[len(family[0]):].strip().removeprefix('-').strip() or 'Base'
 return family,parallel

def build(key,src=None):
 cfg=CONFIGS[key];src=src or read_source(key);rows=src['rows'];groups=OrderedDict();by_number={}
 if len(rows)!=cfg['rows']:raise ValueError('Source row count changed')
 for i,row in enumerate(rows,2):
  family,parallel=classify(key,row[0]);n=row[1];people=tuple(subjects(row[2]));gkey=(family[0],n,people)
  if gkey not in groups:groups[gkey]=[]
  groups[gkey].append((i,row,parallel));by_number.setdefault((family[0],n),set()).add(people)
 cards=[];variants=[];subsets=[];card_for={};counts=Counter()
 for (family,n,people),rr in groups.items():
  name,code,category=next(f for f in cfg['families'] if f[0]==family)[:3];counts[family]+=1
  if len(by_number[(family,n)])>1 and (key,family,n)!=('photogenic','Base Vertical','175'):raise ValueError('Unreviewed conflicting subject')
  canonical=next((r for _,r,p in rr if p=='Base'),rr[0][1]);suffix='-'+slug(canonical[2]) if len(by_number[(family,n)])>1 else ''
  card=dict(id=f"{cfg['prefix']}-{code}-{n}{suffix}",setId=cfg['id'],order=len(cards)+1,number=n,name=' / '.join(clean(s) for s in canonical[2].split('/')),subject2=' / '.join(canonical[2].split('/')[1:]).strip() or None,category=category,subset=family,subsetCode=code,roster=canonical[3] or None,rookie='No',parallelGroup=code)
  cards.append(card)
  for _,r,p in rr:card_for[tuple(r)]=card
 if dict(counts)!={n:c for n,_,_,c in cfg['families']}:raise ValueError('Identity manifest changed '+str(counts))
 seen=set()
 for i,row in enumerate(rows,1):
  card=card_for[tuple(row)];_,parallel=classify(key,row[0]);cap=int(row[4]) if row[4] else None;vid=card['id']+'-'+slug(parallel)
  if vid in seen:raise ValueError('Duplicate source variant '+vid)
  seen.add(vid)
  variants.append(dict(id=vid,set_id=cfg['id'],card_id=card['id'],variant_order=i,subset=card['subset'],card_number=card['number'],display_name=card['name'],parallel=parallel,serial_cap=cap,serial_exact=('1/1' if cap==1 else f'/{cap}') if cap else None,exclusive_note='First Off the Line' if 'FOTL' in parallel else None,numbering_note='Exact listed subject/edition only; blank SEQ is unannounced, not an inferred print run.',verification_status='Published Panini-format checklist; two workbook mirrors reconciled'))
 for n,code,category,count in cfg['families']:subsets.append(dict(category=category,name=n,code=code,count=count))
 record=dict(id=cfg['id'],categoryId='wwe',name=cfg['name'],shortName=cfg['short'],accent=cfg['accent'],year=2024,manufacturer='Panini',releaseDate=cfg['date'],cardCount=len(cards),subsetCount=len(subsets),subsets=subsets)
 return record,cards,variants

def apply(cat,key):
 record,cards,_=build(key);sid=record['id'];old=[s for s in cat['sets'] if s['id']==sid];oc=[c for c in cat['cards'] if c['setId']==sid]
 if old or oc:
  if old!=[record] or oc!=cards:raise ValueError('Existing checklist differs; refusing overwrite')
 else:cat['sets'].append(record);cat['cards'].extend(cards);cat['generatedAt']='2026-10-26T00:00:00Z'
 return cat

def tables(key):
 cfg=CONFIGS[key];r,c,v=build(key);sid=r['id'];out=OrderedDict()
 out['catalogue_sets']=[dict(id=sid,name=r['name'],short_name=r['shortName'],year=2024,manufacturer='Panini',release_date=r['releaseDate'],source_file=cfg['source'],accent=r['accent'],card_count=len(c),subset_count=len(r['subsets']),variant_count=len(v))]
 out['catalogue_subsets']=[dict(set_id=sid,subset_code=s['code'],category=s['category'],name=s['name'],card_count=s['count'],parallel_group=s['code'],source_url=cfg['source']) for s in r['subsets']]
 out['catalogue_cards']=[dict(id=x['id'],set_id=sid,checklist_order=x['order'],category=x['category'],subset_code=x['subsetCode'],card_number=x['number'],display_name=x['name'],subject_1=x['name'].split(' / ')[0],subject_2=x['subject2'],roster=x['roster'],rookie=False,parallel_group=x['parallelGroup'],image_status='Missing',pricing_status='Unpriced',source_url=cfg['source'],notes='Published Panini-format workbook reconciled with an independent workbook mirror. Parallel-only identities retained without inventing absent base issues.'+(' Source #175 differs by edition: base Triple H, parallels Beth Phoenix. Distinct IDs retained; physical numbering unresolved.' if key=='photogenic' and x['subsetCode']=='BV' and x['number']=='175' else '')) for x in c]
 out['catalogue_variants']=v
 rules=OrderedDict();by_id={x['id']:x for x in c}
 for x in v:
  k=(by_id[x['card_id']]['subsetCode'],x['parallel'],x['serial_cap'])
  if k not in rules:
   code,p,cap=k;rules[k]=dict(id=sid+'-'+code.lower()+'-'+slug(p).lower()+'-'+str(cap or 'unannounced'),set_id=sid,parallel_group=code,applies_to=x['subset'],parallel=p,serial_cap=cap,serial_exact=x['serial_exact'],exclusive_note=x['exclusive_note'],numbering_note='Applies ONLY to exact subject/version rows in catalogue_variants; no family-wide grid implied. Similar Red and Blue labels remain separate source labels pending physical verification.',verification_status=x['verification_status'])
 out['catalogue_parallel_rules']=list(rules.values())
 out['catalogue_sources']=[dict(id=sid+'-'+suffix,set_id=sid,source=source,purpose=purpose,url=url,notes=note) for suffix,source,purpose,url,note in [
  ('checklist','Panini-format checklist / Checklist Insider mirror','Checklist',cfg['source'],'Full published workbook; source rows retained, two mirrors reconciled. Original GTS workbook also reconciled for Three Count/Select after one documented Seth nickname encoding repair. PhotoGenic direct-origin workbook not recovered.'),
  ('comparison','Checklist Center','Independent workbook comparison',cfg['center'],'All rows match after numeric normalization, signer-order and repeated Progressions subject normalization, and combining dual signer rows.'),
  ('release','Checklist Insider','Release date',cfg['guide'],'Reported final release date; Three Count GTS preliminary Nov 22 schedule differs from Dec 6 final guide date.'),
  ('manufacturer','Panini America','Manufacturer checklist index','https://www.paniniamerica.net/checklist.html','Manufacturer index retained; current product pages no longer provide complete historical content.')]]
 src=read_source(key)
 if src.get('distributorUrl'):
  out['catalogue_sources'].append(dict(id=sid+'-distributor',set_id=sid,source='GTS Distribution / Panini',purpose='Original distributor checklist',url=src['distributorUrl'],notes=f"All {src['distributorRows']} rows reconciled. {src['encodingRepairs']} instances of one Seth nickname mojibake string repaired only for comparison/display. Original rows and workbook hash retained."))
 return out

if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--extract',action='store_true');a=p.parse_args()
 if a.extract:
  for k in CONFIGS:
   src=extract(k);(SOURCE_DIR/(k+'.json')).write_text(json.dumps(src,ensure_ascii=False,separators=(',',':'))+'\n');print(k,'independently reconciled',len(src['rows']),'rows')
 else:
  cat=load_catalogue()
  for k in CONFIGS:
   apply(cat,k);r,c,v=build(k);print(k,len(c),'identities',len(r['subsets']),'subsets',len(v),'versions')
  save_catalogue(cat)
